"""Turn scraped rows into downloadable files: Excel, CSV, PDF, JSON and plain text."""
import csv
import io
import json
from datetime import datetime

from fpdf import FPDF, FontFace
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

FONT_DIR = "/usr/share/fonts/truetype/dejavu"


def to_excel(rows, fields, title):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append([f.replace("_", " ").title() for f in fields])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F4A3A")
    for row in rows:
        sheet.append([row.get(f, "") for f in fields])
    # Make each column wide enough to read (up to a limit), and keep the header visible
    for i, field in enumerate(fields, start=1):
        longest = max([len(field)] + [len(str(r.get(field, ""))) for r in rows])
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = min(longest + 3, 60)
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def to_csv(rows, fields, title):
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fields)
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")  # the "sig" helps Excel show symbols like £ correctly


def to_json(rows, fields, title):
    return json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8")


def to_text(rows, fields, title):
    widths = {f: min(max([len(f)] + [len(str(r.get(f, ""))) for r in rows]), 50) for f in fields}
    line = lambda values: "  ".join(str(v)[:widths[f]].ljust(widths[f]) for f, v in zip(fields, values))
    lines = [title, "", line([f.upper() for f in fields]), line(["-" * widths[f] for f in fields])]
    lines += [line([r.get(f, "") for f in fields]) for r in rows]
    return "\n".join(lines).encode("utf-8")


def to_pdf(rows, fields, title):
    pdf = FPDF(orientation="landscape")
    pdf.add_font("DejaVu", "", f"{FONT_DIR}/DejaVuSans.ttf")
    pdf.add_font("DejaVu", "B", f"{FONT_DIR}/DejaVuSans-Bold.ttf")
    pdf.add_font("Serif", "B", f"{FONT_DIR}/DejaVuSerif-Bold.ttf")
    pdf.set_auto_page_break(True, margin=12)
    pdf.add_page()
    pdf.set_text_color(47, 74, 58)
    pdf.set_font("Serif", "B", 18)
    pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(107, 103, 92)
    pdf.set_font("DejaVu", "", 9)
    pdf.cell(0, 6, f"{len(rows):,} records  •  collected {datetime.now():%d %b %Y, %H:%M}  •  Solvix",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_text_color(31, 42, 36)
    pdf.set_font("DejaVu", "", 8)
    heading_style = FontFace(emphasis="BOLD", color=(244, 239, 230), fill_color=(47, 74, 58))
    # Give long columns (like titles) more room than short ones (like ratings)
    widths = [min(max([len(f)] + [len(str(r.get(f, ""))) for r in rows]), 60) + 4 for f in fields]
    with pdf.table(text_align="LEFT", line_height=5, headings_style=heading_style, col_widths=widths,
                   borders_layout="NONE", padding=1.5,
                   cell_fill_color=(244, 239, 230), cell_fill_mode="ROWS") as table:
        table.row([f.replace("_", " ").upper() for f in fields])
        for r in rows:
            table.row([str(r.get(f, "")) for f in fields])
    return bytes(pdf.output())


# file extension: (function that makes the file, file type for the browser)
FORMATS = {
    "xlsx": (to_excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "csv":  (to_csv, "text/csv"),
    "pdf":  (to_pdf, "application/pdf"),
    "json": (to_json, "application/json"),
    "txt":  (to_text, "text/plain"),
}
